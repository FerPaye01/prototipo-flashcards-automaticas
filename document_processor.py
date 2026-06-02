import os
import fitz  # PyMuPDF
from docx import Document
from typing import List, Dict, Any

class DocumentProcessor:
    """
    Procesador de documentos (PDF, Word) para extraer texto por páginas o párrafos.
    """
    
    def __init__(self, log_callback=None):
        self.log_callback = log_callback

    def _log(self, message):
        if self.log_callback:
            self.log_callback(message)
        else:
            print(message)

    def render_pdf_pages_to_images(self, pdf_path: str, output_dir: str) -> List[Dict[str, Any]]:
        """
        Renderiza cada página de un PDF como imagen y las guarda en output_dir.
        """
        pages_content = []
        try:
            if not os.path.exists(output_dir):
                os.makedirs(output_dir, exist_ok=True)
                
            doc = fitz.open(pdf_path)
            self._log(f"   🖼️ Renderizando PDF a imágenes: {os.path.basename(pdf_path)} ({len(doc)} páginas)")
            
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                # Zoom 2x para mejor calidad de OCR posterior
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                
                image_name = f"page_{page_num + 1}.png"
                image_path = os.path.join(output_dir, image_name)
                pix.save(image_path)
                
                pages_content.append({
                    "page_number": page_num + 1,
                    "image_path": image_path,
                    "content": page.get_text().strip()
                })
            doc.close()
            return pages_content
        except Exception as e:
            self._log(f"   ❌ Error al renderizar PDF: {e}")
            return []

    def split_pdf(self, pdf_path: str, page_numbers: List[int], output_path: str) -> bool:
        """
        Crea un nuevo PDF con las páginas especificadas (1-indexed).
        """
        try:
            doc = fitz.open(pdf_path)
            new_doc = fitz.open()
            for p_num in page_numbers:
                # fitz usa índices basados en 0
                new_doc.insert_pdf(doc, from_page=p_num-1, to_page=p_num-1)
            
            # Asegurar que el directorio de salida exista
            os.makedirs(os.path.dirname(output_path), exist_ok=True)
            
            new_doc.save(output_path)
            new_doc.close()
            doc.close()
            return True
        except Exception as e:
            self._log(f"   ❌ Error al dividir PDF: {e}")
            return False

    def extract_text_from_pdf(self, pdf_path: str) -> List[Dict[str, Any]]:
        """
        Extrae texto de un PDF página por página.
        """
        pages_content = []
        try:
            doc = fitz.open(pdf_path)
            self._log(f"   📖 PDF abierto: {os.path.basename(pdf_path)} ({len(doc)} páginas)")
            
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                text = page.get_text().strip()
                pages_content.append({
                    "page_number": page_num + 1,
                    "content": text
                })
            doc.close()
            return pages_content
        except Exception as e:
            self._log(f"   ❌ Error al extraer PDF: {e}")
            return []

    def extract_text_from_docx(self, docx_path: str) -> List[Dict[str, Any]]:
        """
        Extrae texto de un archivo Word (.docx). 
        Como Word no tiene páginas fijas, dividimos por párrafos o grupos de párrafos.
        """
        pages_content = []
        try:
            doc = Document(docx_path)
            self._log(f"   📖 Word abierto: {os.path.basename(docx_path)} ({len(doc.paragraphs)} párrafos)")
            
            # Agrupamos párrafos para simular 'páginas' (aprox 500 palabras por 'página')
            current_page_text = []
            word_count = 0
            page_num = 1
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                    
                current_page_text.append(text)
                word_count += len(text.split())
                
                if word_count >= 500:
                    pages_content.append({
                        "page_number": page_num,
                        "content": "\n".join(current_page_text)
                    })
                    current_page_text = []
                    word_count = 0
                    page_num += 1
            
            # Añadir última página si hay contenido
            if current_page_text:
                pages_content.append({
                    "page_number": page_num,
                    "content": "\n".join(current_page_text)
                })
                
            return pages_content
        except Exception as e:
            self._log(f"   ❌ Error al extraer Word: {e}")
            return []

    def segment_pages(self, pages: List[Dict[str, Any]], pages_per_section: int, overlap: int) -> List[Dict[str, Any]]:
        """
        Agrupa páginas en secciones con un solape determinado.
        """
        sections = []
        total_pages = len(pages)
        start = 0
        
        while start < total_pages:
            end = min(start + pages_per_section, total_pages)
            group = pages[start:end]
            
            section_content = ""
            page_range = f"Páginas {group[0]['page_number']}-{group[-1]['page_number']}"
            
            combined_text = []
            images = []
            for pg in group:
                combined_text.append(f"--- Página {pg['page_number']} ---\n{pg['content']}")
                if "image_path" in pg:
                    images.append(pg["image_path"])
            
            section_content = "\n\n".join(combined_text)
            
            sections.append({
                "title": f"Sección ({page_range})",
                "content": section_content,
                "images": images,
                "pages": [pg['page_number'] for pg in group]
            })
            
            if end >= total_pages:
                break
                
            # Avanzar respetando el solape
            start = end - overlap
            if start <= (end - pages_per_section):
                start = end + 1
                
        return sections

    def extract_full_text_with_page_markers(self, pdf_path: str) -> str:
        """
        Extrae el texto completo de un PDF con marcadores de página.
        Formato: --- PÁGINA X --- seguido del texto de esa página.
        Útil para que un LLM pueda referenciar páginas específicas.
        
        Returns:
            Texto completo con marcadores, o cadena vacía si falla.
        """
        try:
            doc = fitz.open(pdf_path)
            total_pages = len(doc)
            self._log(f"   📖 Extrayendo texto con marcadores de página: {os.path.basename(pdf_path)} ({total_pages} páginas)")
            
            parts = []
            for page_num in range(total_pages):
                page = doc.load_page(page_num)
                text = page.get_text().strip()
                if text:
                    parts.append(f"--- PÁGINA {page_num + 1} ---\n{text}")
                else:
                    parts.append(f"--- PÁGINA {page_num + 1} ---\n[Página sin texto extraíble - posiblemente contiene solo imágenes o gráficos]")
            
            doc.close()
            full_text = "\n\n".join(parts)
            self._log(f"   ✅ {total_pages} páginas extraídas ({len(full_text):,} caracteres)")
            return full_text
            
        except Exception as e:
            self._log(f"   ❌ Error extrayendo texto con marcadores: {e}")
            return ""

    def extract_text_per_page(self, pdf_path: str) -> List[Dict[str, Any]]:
        """
        Extrae el texto de cada página individualmente.
        Útil para embeddings por página (segmentación semántica local con BGE-M3).
        
        Returns:
            Lista de dicts con page_number y text.
        """
        pages = []
        try:
            doc = fitz.open(pdf_path)
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                text = page.get_text().strip()
                pages.append({
                    "page_number": page_num + 1,
                    "text": text
                })
            doc.close()
            return pages
        except Exception as e:
            self._log(f"   ❌ Error extrayendo texto por página: {e}")
            return []
